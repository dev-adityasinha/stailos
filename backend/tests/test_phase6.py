"""Phase 6 verification: analytics, reports (csv/xlsx/pdf), documents."""
import io

from tests.conftest import auth_headers, create_user_as_admin
from tests.test_leads import make_lead
from tests.test_phase3 import make_customer, make_project


def _seed_sales_data(client, admin):
    """One lead→customer→booking with a payment, plus a lost lead."""
    make_project(client, admin)
    units = client.get("/api/v1/properties", headers=admin).json()["data"]
    lead = make_lead(client, admin).json()["data"]
    customer = client.post(
        f"/api/v1/leads/{lead['id']}/convert", headers=admin, json={}
    ).json()["data"]
    booking = client.post(
        "/api/v1/bookings", headers=admin,
        json={"customer_id": customer["id"], "unit_id": units[0]["id"],
              "total_value": 7500000, "lead_id": lead["id"]},
    ).json()["data"]
    client.post(
        f"/api/v1/bookings/{booking['id']}/payments", headers=admin,
        json={"amount": 1500000, "milestone": "Token + down payment"},
    )
    lost = make_lead(client, admin, full_name="Lost Person", phone="9111111111",
                     email="lost@x.com", source="meta_ads", force=True).json()["data"]
    client.patch(
        f"/api/v1/pipeline/leads/{lost['id']}/stage", headers=admin,
        json={"stage": "lost", "lost_reason": "Budget"},
    )
    return lead, customer, booking, units


class TestAnalytics:
    def test_overview_counts(self, client, admin):
        _seed_sales_data(client, admin)
        data = client.get("/api/v1/analytics/overview", headers=admin).json()["data"]
        assert data["total_leads"] == 2
        assert data["active_leads"] == 1  # one lost
        assert data["converted_leads"] == 1
        assert data["conversion_rate"] == 50.0
        assert data["bookings_total"] == 1
        assert data["revenue_collected"] == 1500000.0
        assert data["pipeline_value"] == 7500000.0

    def test_funnel_has_all_stages(self, client, admin):
        _seed_sales_data(client, admin)
        funnel = client.get("/api/v1/analytics/funnel", headers=admin).json()["data"]
        assert len(funnel) == 9
        by_stage = {r["stage"]: r["count"] for r in funnel}
        assert by_stage["lost"] == 1

    def test_sources_and_team(self, client, admin):
        _seed_sales_data(client, admin)
        sources = client.get("/api/v1/analytics/sources", headers=admin).json()["data"]
        by_source = {r["source"]: r for r in sources}
        assert by_source["website"]["converted"] == 1
        assert by_source["meta_ads"]["leads"] == 1
        team = client.get("/api/v1/analytics/team", headers=admin).json()["data"]
        assert team[0]["revenue"] == 1500000.0

    def test_revenue_series(self, client, admin):
        _seed_sales_data(client, admin)
        series = client.get("/api/v1/analytics/revenue?months=3", headers=admin).json()["data"]
        assert len(series) == 3
        assert series[-1]["revenue"] == 1500000.0  # current month

    def test_scope_limits_analytics(self, client, admin):
        _seed_sales_data(client, admin)
        create_user_as_admin(client, admin, "rep@stail.com", "sales_executive")
        rep = auth_headers(client, "rep@stail.com")
        data = client.get("/api/v1/analytics/overview", headers=rep).json()["data"]
        assert data["total_leads"] == 0 and data["revenue_collected"] == 0.0


class TestReports:
    def test_catalog(self, client, admin):
        data = client.get("/api/v1/reports/catalog", headers=admin).json()["data"]
        assert set(data["formats"]) == {"csv", "xlsx", "pdf"}
        assert "revenue" in data["types"]

    def test_csv_report(self, client, admin):
        _seed_sales_data(client, admin)
        resp = client.post(
            "/api/v1/reports/generate", headers=admin,
            json={"type": "lead", "format": "csv"},
        )
        assert resp.status_code == 200
        assert resp.headers["x-row-count"] == "2"
        assert "Rahul Sharma" in resp.text

    def test_xlsx_report(self, client, admin):
        _seed_sales_data(client, admin)
        resp = client.post(
            "/api/v1/reports/generate", headers=admin,
            json={"type": "sales", "format": "xlsx"},
        )
        assert resp.status_code == 200
        assert resp.content[:2] == b"PK"  # zip container
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        assert ws.cell(row=1, column=1).value == "Customer"
        assert ws.max_row == 2

    def test_pdf_report(self, client, admin):
        _seed_sales_data(client, admin)
        resp = client.post(
            "/api/v1/reports/generate", headers=admin,
            json={"type": "revenue", "format": "pdf"},
        )
        assert resp.status_code == 200
        assert resp.content[:5] == b"%PDF-"

    def test_all_report_types_generate(self, client, admin):
        _seed_sales_data(client, admin)
        for report_type in ["lead", "sales", "agent", "booking", "property", "revenue"]:
            resp = client.post(
                "/api/v1/reports/generate", headers=admin,
                json={"type": report_type, "format": "csv"},
            )
            assert resp.status_code == 200, report_type

    def test_invalid_type_rejected(self, client, admin):
        resp = client.post(
            "/api/v1/reports/generate", headers=admin,
            json={"type": "nonsense", "format": "csv"},
        )
        assert resp.status_code == 400


class TestDocuments:
    def _upload(self, client, headers, **form):
        data = {"title": "Booking agreement", "category": "agreement"}
        data.update(form)
        return client.post(
            "/api/v1/documents/upload", headers=headers, data=data,
            files={"file": ("agreement.pdf", io.BytesIO(b"%PDF-1.4 v1"), "application/pdf")},
        )

    def test_upload_list_download(self, client, admin):
        resp = self._upload(client, admin)
        assert resp.status_code == 201, resp.text
        doc = resp.json()["data"]
        assert doc["current_version"] == 1

        listing = client.get("/api/v1/documents", headers=admin).json()
        assert listing["meta"]["total"] == 1

        dl = client.get(f"/api/v1/documents/{doc['id']}/download", headers=admin)
        assert dl.status_code == 200
        assert dl.content == b"%PDF-1.4 v1"

    def test_version_tracking(self, client, admin):
        doc = self._upload(client, admin).json()["data"]
        resp = client.post(
            f"/api/v1/documents/{doc['id']}/version", headers=admin,
            data={"note": "Signed copy"},
            files={"file": ("agreement_signed.pdf", io.BytesIO(b"%PDF-1.4 v2"),
                            "application/pdf")},
        )
        assert resp.status_code == 201
        updated = resp.json()["data"]
        assert updated["current_version"] == 2
        assert len(updated["versions"]) == 2
        # Latest by default, older on request.
        latest = client.get(f"/api/v1/documents/{doc['id']}/download", headers=admin)
        assert latest.content == b"%PDF-1.4 v2"
        old = client.get(
            f"/api/v1/documents/{doc['id']}/download?version=1", headers=admin
        )
        assert old.content == b"%PDF-1.4 v1"

    def test_entity_attachment_and_timeline(self, client, admin):
        customer = make_customer(client, admin).json()["data"]
        resp = self._upload(
            client, admin, title="KYC PAN card", category="kyc",
            entity_type="customer", entity_id=customer["id"],
        )
        assert resp.status_code == 201
        docs = client.get(
            f"/api/v1/documents?entity_type=customer&entity_id={customer['id']}",
            headers=admin,
        ).json()
        assert docs["meta"]["total"] == 1
        full = client.get(
            f"/api/v1/customers/{customer['id']}/360", headers=admin
        ).json()["data"]
        assert any(a["type"] == "document" for a in full["timeline"])

    def test_scope_own_isolation(self, client, admin):
        create_user_as_admin(client, admin, "repa@stail.com", "sales_executive")
        create_user_as_admin(client, admin, "repb@stail.com", "sales_executive")
        rep_a, rep_b = auth_headers(client, "repa@stail.com"), auth_headers(client, "repb@stail.com")
        doc = self._upload(client, rep_a).json()["data"]
        assert client.get("/api/v1/documents", headers=rep_b).json()["meta"]["total"] == 0
        assert client.get(f"/api/v1/documents/{doc['id']}", headers=rep_b).status_code == 404

    def test_dangerous_extension_rejected(self, client, admin):
        resp = client.post(
            "/api/v1/documents/upload", headers=admin,
            data={"title": "Nasty file"},
            files={"file": ("evil.sh", io.BytesIO(b"#!/bin/sh"), "text/x-sh")},
        )
        assert resp.status_code == 400
